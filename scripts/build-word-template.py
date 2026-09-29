"""Builds public/templates/Mau_soan_de_Word.docx, the Word template for the question import.

Run from the repo root: python scripts/build-word-template.py  (needs python-docx). The questions are examples
written for the template, not taken from an exam. src/modules/question-bank/domain/word-template.test.ts reads the
file back through the app's reader, so change both together.
"""
import io
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

RED = RGBColor(0xFF, 0x00, 0x00)
BLUE = RGBColor(0x28, 0x6F, 0xB7)
OUT = 'public/templates/Mau_soan_de_Word.docx'


def warning_sign_png(width=180, height=160):
    """A yellow warning triangle with a black border, drawn pixel by pixel (no imaging library needed)."""
    rows = []
    top, bottom = 10, height - 10
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            inside = outer = False
            if top <= y <= bottom:
                half = (y - top) / (bottom - top) * (width / 2 - 10)
                outer = abs(x - width / 2) <= half
                inner_half = half - 12
                inside = abs(x - width / 2) <= inner_half and y <= bottom - 10
            mark = inside and abs(x - width / 2) <= 6 and (60 <= y <= 110 or 120 <= y <= 132)
            if mark or (outer and not inside):
                row += bytes([20, 20, 20])
            elif inside:
                row += bytes([255, 204, 0])
            else:
                row += bytes([255, 255, 255])
        rows.append(bytes(row))
    raw = b''.join(rows)

    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(raw, 9)) + chunk(b'IEND', b''))


def para(document, text='', bold=False, color=None, size=None):
    paragraph = document.add_paragraph()
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        if color is not None:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)
    return paragraph


def question(document, number, stem, tags=''):
    paragraph = document.add_paragraph()
    head = paragraph.add_run(f'Câu {number}. ')
    head.bold = True
    if tags:
        tag = paragraph.add_run(tags + ' ')
        tag.bold = True
        tag.font.color.rgb = BLUE
    paragraph.add_run(stem)
    return paragraph


def options(document, items):
    """items: (text, correct) pairs, typed as a), b), c)…; correct ones in red."""
    for index, (text, correct) in enumerate(items):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f'{chr(97 + index)}) {text}')
        if correct:
            run.font.color.rgb = RED


def main():
    document = Document()
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(13)

    title = para(document, 'MẪU SOẠN ĐỀ TRẮC NGHIỆM ĐỂ NHẬP VÀO APP THI ONLINE HPTTS', bold=True, size=14)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(document, 'Phần hướng dẫn này nằm trước "Câu 1" nên app bỏ qua; giữ lại hay xóa đều được.')
    for line in [
        '1. Mỗi câu bắt đầu ở đầu dòng bằng chữ "Câu", số thứ tự và dấu chấm hoặc hai chấm. Không gõ nhầm "Cây", "Cau".',
        '2. Mỗi phương án một dòng: a), b), c), d)… (gõ tay hoặc dùng đánh số tự động của Word). Tối đa 10 phương án.',
        '3. Đáp án đúng: tô ĐỎ cả dòng phương án. Không tô đỏ chữ khác trong câu.',
        '4. Dạng câu khác trắc nghiệm một đáp án: ghi nhãn ngay sau "Câu N.": [Nhiều đáp án], [Đúng/Sai], [Nối cột], [Sắp xếp], [Tự luận].',
        '5. Nhãn thêm (không bắt buộc): [Dễ], [Trung bình], [Khó], [3 điểm]. Không ghi thì là Trung bình, 2 điểm.',
        '6. [Đúng/Sai]: tô đỏ các ý Đúng, ý để đen là Sai.',
        '7. [Nối cột]: bảng 2 cột ngay dưới đề bài, mỗi hàng là một cặp đúng. Khi thi, app tự xáo cột phải.',
        '8. [Sắp xếp]: ghi các bước theo đúng thứ tự. Khi thi, app tự xáo.',
        '9. [Tự luận]: mỗi ý chấm một dòng "Ý chấm: nội dung | điểm" (thiếu điểm thì tính 2).',
        '10. Ảnh: Insert → Pictures, đặt ngay dưới đề bài, chọn "In Line with Text"; mỗi câu một ảnh, dạng PNG hoặc JPG.',
        '11. Không cần: dòng "Đáp án: B" (nếu có thì phải trùng với phương án tô đỏ), phần "Giải thích:" (app bỏ qua).',
        '12. Lưu file dạng .docx (Word 2007 trở lên). File .doc cũ: File → Save As → Word Document.',
    ]:
        para(document, line)
    para(document)

    question(document, 1, 'Khi nâng hàng bằng cần trục, vị trí nào nguy hiểm nhất?')
    options(document, [('Phía dưới tải trọng đang nâng', True), ('Bên cạnh cabin điều khiển', False),
                       ('Ngoài khu vực làm việc', False), ('Trên sàn công tác', False)])

    question(document, 2, 'Thiết bị nào sau đây thuộc nhóm thiết bị nâng?', '[Nhiều đáp án]')
    options(document, [('Cầu trục', True), ('Palăng xích', True), ('Máy bơm nước', False), ('Xe đẩy tay', False)])

    question(document, 3, 'Xác định Đúng hoặc Sai cho từng phát biểu về an toàn vận hành:', '[Đúng/Sai] [3 điểm]')
    options(document, [('Phải kiểm tra khu vực làm việc trước khi nâng hàng.', True),
                       ('Được phép nâng vượt tải 20% khi khẩn cấp.', False),
                       ('Các thiết bị hạn vị phải hoạt động trước mỗi ca.', True),
                       ('Không cần ngắt nguồn khi bảo trì nhỏ.', False)])

    question(document, 4, 'Nối thiết bị với chức năng đúng:', '[Nối cột]')
    table = document.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for left, right in [('Bộ chống lắc hàng', 'Giảm dao động của hàng khi xe con tăng, giảm tốc'),
                        ('Công tắc hành trình', 'Ngắt mạch khi cơ cấu chạm điểm giới hạn'),
                        ('Thiết bị đo tải', 'Ngắt tời nâng khi quá tải')]:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right

    question(document, 5, 'Sắp xếp các bước nâng hàng theo đúng thứ tự:', '[Sắp xếp] [Dễ]')
    options(document, [('Kiểm tra tải trọng và dụng cụ mang tải', False), ('Móc cáp vào hàng', False),
                       ('Quan sát vùng nguy hiểm', False), ('Ra tín hiệu nâng', False)])

    question(document, 6, 'Nêu các nguyên nhân thường gặp gây tai nạn lao động khi xếp dỡ hàng.', '[Tự luận] [Khó]')
    for line in ['Ý chấm: làm sai quy trình | 2', 'Ý chấm: thiếu trang bị bảo hộ | 2', 'Ý chấm: không kiểm tra thiết bị trước ca | 2']:
        para(document, line)

    question(document, 7, 'Biển báo trong hình dưới đây có ý nghĩa gì?')
    document.add_paragraph().add_run().add_picture(io.BytesIO(warning_sign_png()), width=Cm(3))
    options(document, [('Cấm đi qua', False), ('Chú ý nguy hiểm', True), ('Lối thoát hiểm', False), ('Khu vực hút thuốc', False)])

    document.save(OUT)
    print('Built', OUT)


if __name__ == '__main__':
    main()
